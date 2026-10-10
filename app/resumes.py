"""Resume intake with exact-source facts, never inferred candidate achievements."""

from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from typing import Literal

from docx import Document
from pydantic import BaseModel, ConfigDict, Field
from pypdf import PdfReader
from sqlmodel import Session

from app.db import get_profile
from app.llm import call_tool
from app.models import ResumeProfile, utcnow


class ExtractedFact(BaseModel):
    model_config = ConfigDict(extra="forbid")
    category: Literal["skill", "experience", "education", "project", "identity"]
    evidence: str = Field(min_length=1)


class ResumeExtraction(BaseModel):
    model_config = ConfigDict(extra="forbid")
    facts: list[ExtractedFact] = Field(min_length=1)


def latex_text(source: str) -> str:
    text = re.sub(r"(?<!\\)%[^\n]*", "", source)
    text = re.sub(r"\\(?:href|url)\{[^{}]*\}", "", text)
    text = re.sub(r"\\(?:documentclass|usepackage|begin|end|label|pagestyle|setlength)(?:\[[^\]]*\])?\{[^{}]*\}", "", text)
    text = re.sub(r"\\[a-zA-Z@]+\*?(?:\[[^\]]*\])?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    text = re.sub(r"\\([%&_#])", r"\1", text)
    return "\n".join(re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip())


def read_upload(filename: str, content: bytes) -> tuple[str, str | None]:
    ext = Path(filename).suffix.lower()
    try:
        if ext in (".txt", ".tex"):
            source = content.decode("utf-8")
            text, latex = (latex_text(source), source) if ext == ".tex" else (source, None)
        elif ext == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Upload an unencrypted PDF")
            text, latex = "\n".join(page.extract_text() or "" for page in reader.pages), None
        elif ext == ".docx":
            doc = Document(BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
            latex = None
        else:
            raise ValueError("Supported resume formats: .tex, .pdf, .docx, .txt")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Cannot read this resume; upload a valid, text-readable file") from exc
    if not text.strip():
        raise ValueError("No readable text found. Scanned PDFs require a text-readable version or .tex source")
    return text.strip(), latex


def grounded_facts(extraction: ResumeExtraction, text: str) -> dict:
    facts, seen = [], set()
    for fact in extraction.facts:
        if fact.evidence not in text:
            raise ValueError("Resume extraction contained a claim without exact source evidence")
        key = (fact.category, fact.evidence)
        if key in seen:
            continue
        seen.add(key)
        facts.append({"id": "fact-" + sha256(fact.evidence.encode()).hexdigest()[:16],
                      "category": fact.category, "text": fact.evidence, "evidence": fact.evidence})
    return {"facts": facts}


async def extract_profile(text: str) -> dict:
    result = await call_tool(ResumeExtraction, system=(
        "Extract independently useful resume facts. Every evidence field must be an exact, contiguous quote "
        "from resume_text. Preserve whole achievement sentences/bullets, skills, education, roles, and project "
        "statements. Do not infer missing facts, employers, metrics, dates or skills. Assign only a category."
    ), data={"resume_text": text})
    return grounded_facts(result, text)


def save_profile(session: Session, *, filename: str, text: str, latex: str | None, facts: dict):
    profile = get_profile(session)
    if profile is None:
        profile = ResumeProfile(user_id=1, filename=Path(filename).name, raw_resume_text=text)
    profile.filename = Path(filename).name
    profile.raw_resume_text = text
    profile.latex_source = latex
    profile.structured_facts = facts
    profile.updated_at = utcnow()
    session.add(profile)
    session.commit()
    session.refresh(profile)
    return profile
